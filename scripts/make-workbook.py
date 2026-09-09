"""
Build the sample-data workbook from the exported dataset.

    node scripts/export-dataset.mjs > dataset.json
    python3 scripts/make-workbook.py dataset.json AskData-sample-data.xlsx

One sheet per table, plus a Companies sheet first - that one carries the
per-service org ids, which are what someone testing the app actually needs to
type in. PII columns are headed in red so a reader can see at a glance which
columns AskData masks on the way out.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SRC = sys.argv[1] if len(sys.argv) > 1 else "dataset.json"
OUT = sys.argv[2] if len(sys.argv) > 2 else "AskData-sample-data.xlsx"

data = json.load(open(SRC))

HEAD_FILL = PatternFill("solid", fgColor="1F3864")
PII_FILL = PatternFill("solid", fgColor="8B2A2A")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=10)
TITLE_FONT = Font(bold=True, size=13)
NOTE_FONT = Font(italic=True, size=9, color="666666")

wb = Workbook()


def style_header(ws, columns, pii_columns, row=1):
    for i, name in enumerate(columns, start=1):
        cell = ws.cell(row=row, column=i, value=name)
        cell.font = HEAD_FONT
        cell.fill = PII_FILL if name in pii_columns else HEAD_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def autosize(ws, columns, rows, max_width=46):
    for i, name in enumerate(columns, start=1):
        widest = len(str(name))
        for row in rows[:200]:
            value = row.get(name)
            if value is not None:
                widest = max(widest, len(str(value)))
        ws.column_dimensions[get_column_letter(i)].width = min(widest + 3, max_width)


# ---------------------------------------------------------------- companies
ws = wb.active
ws.title = "Companies"
ws["A1"] = "AskData sample customers - 10 companies"
ws["A1"].font = TITLE_FONT
ws["A2"] = (
    "Pick a Service in AskData, then the org id for that service from this sheet. "
    "A dash means the company is not subscribed to that service - AskData refuses "
    "those questions rather than returning an empty answer."
)
ws["A2"].font = NOTE_FONT

COMPANY_COLS = [
    "ORG_NAME", "ORG_ID", "SUBSCRIBED_PRODUCTS", "CRM_ORG_ID", "CMP_ORG_ID",
    "DESK_ORG_ID", "OPEN_TICKET", "DC", "EDITION", "SIGNED_UP_ON",
]
style_header(ws, COMPANY_COLS, set(), row=4)
ws.freeze_panes = "A5"
for r, company in enumerate(data["companies"], start=5):
    for c, name in enumerate(COMPANY_COLS, start=1):
        ws.cell(row=r, column=c, value=company.get(name) or "—")
autosize(ws, COMPANY_COLS, data["companies"])

# -------------------------------------------------------------- one per table
SERVICE_ORDER = {"CRM_": 1, "CMP_": 2, "DESK_": 3}


def sort_key(name):
    for prefix, rank in SERVICE_ORDER.items():
        if name.startswith(prefix):
            return (rank, name)
    return (0, name)


summary = []
for name in sorted(data["tables"], key=sort_key):
    rows = data["tables"][name]
    columns = data["columns"].get(name) or (list(rows[0]) if rows else [])
    pii = set(data["pii"].get(name, []))

    ws = wb.create_sheet(name[:31])
    style_header(ws, columns, pii)
    for r, row in enumerate(rows, start=2):
        for c, column in enumerate(columns, start=1):
            value = row.get(column)
            if isinstance(value, str) and value in ("true", "false"):
                value = value == "true"
            ws.cell(row=r, column=c, value=value)
    autosize(ws, columns, rows)
    summary.append((name, len(rows), len(columns), ", ".join(sorted(pii)) or "—"))

# ------------------------------------------------------------------ contents
ws = wb.create_sheet("Contents", 1)
ws["A1"] = "Tables in this workbook"
ws["A1"].font = TITLE_FONT
ws["A2"] = "Columns shaded red hold PII. AskData masks those server-side; revealing one row is audited."
ws["A2"].font = NOTE_FONT
CONTENTS_COLS = ["TABLE", "ROWS", "COLUMNS", "PII COLUMNS"]
style_header(ws, CONTENTS_COLS, set(), row=4)
for r, entry in enumerate(summary, start=5):
    for c, value in enumerate(entry, start=1):
        ws.cell(row=r, column=c, value=value)
ws.cell(row=len(summary) + 5, column=1, value="TOTAL").font = Font(bold=True)
ws.cell(row=len(summary) + 5, column=2, value=sum(s[1] for s in summary)).font = Font(bold=True)
autosize(ws, CONTENTS_COLS, [dict(zip(CONTENTS_COLS, s)) for s in summary])

wb.save(OUT)
print(f"{OUT}: {len(wb.sheetnames)} sheets, {sum(s[1] for s in summary)} data rows")
